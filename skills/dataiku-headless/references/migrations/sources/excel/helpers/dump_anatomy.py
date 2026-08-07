#!/usr/bin/env python3
"""workbook anatomy dump · in: <workbook.xlsx> [--json | --schema SHEET [--header-row N]] → out: stdout structure report or Dataiku schema JSON · deps: openpyxl

Reports, per sheet: dimensions, tables, pivot tables, charts, distinct formula
shapes with counts, merged cells, and autofilters. Also reports workbook-level
defined names, external connections, and embedded Power Query source when
present.

Usage:
    python dump_anatomy.py BOOK.xlsx [--json] [--max-formulas N]
    python dump_anatomy.py BOOK.xlsx --schema "Sheet Name"
"""

from __future__ import annotations

import argparse
import base64
import io
import json
import re
import sys
import zipfile
from collections import Counter
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def col_of(cell_ref: str) -> str:
    m = re.match(r"([A-Z]+)", cell_ref)
    return m.group(1) if m else ""


def normalize_formula(f: str) -> str:
    """Collapse row numbers so per-row copies of one formula dedupe together."""
    return re.sub(r"(?<![A-Za-z0-9_])(\$?[A-Z]{1,3})\$?\d+", r"\1#", f)


def extract_m_queries(path: str) -> list[dict]:
    """Pull Power Query section files out of DataMashup customXml parts."""
    out = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not re.match(r"customXml/item\d+\.xml$", name):
                continue
            raw = z.read(name)
            enc = "utf-16" if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8"
            data = raw.decode(enc, "replace")
            m = re.search(r"<DataMashup[^>]*>(.*?)</DataMashup>", data, re.S)
            if not m:
                continue
            try:
                blob = base64.b64decode(m.group(1))
            except Exception:
                continue
            if len(blob) < 8:
                continue
            pkg_len = int.from_bytes(blob[4:8], "little")
            pkg = blob[8 : 8 + pkg_len]
            if not pkg.startswith(b"PK\x03\x04"):
                pk = blob.find(b"PK\x03\x04")
                if pk < 0:
                    continue
                pkg = blob[pk:]
            try:
                inner = zipfile.ZipFile(io.BytesIO(pkg))
            except Exception:
                continue
            for iname in inner.namelist():
                if iname.endswith(".m"):
                    out.append(
                        {
                            "part": name,
                            "file": iname,
                            "code": inner.read(iname).decode("utf-8", "replace"),
                        }
                    )
    return out


def extract_connections(path: str) -> list[dict]:
    """Summarize xl/connections.xml external sources."""
    import xml.etree.ElementTree as ET

    out = []
    with zipfile.ZipFile(path) as z:
        if "xl/connections.xml" not in z.namelist():
            return out
        root = ET.fromstring(z.read("xl/connections.xml"))
        ns = {"m": root.tag.split("}")[0].strip("{")} if "}" in root.tag else {}
        conns = root.findall("m:connection", ns) if ns else root.findall("connection")
        for c in conns:
            info = {
                "name": c.get("name"),
                "description": c.get("description"),
                "type": c.get("type"),
            }
            db = c.find("m:dbPr", ns) if ns else c.find("dbPr")
            if db is not None:
                info["connection"] = db.get("connection")
                info["command"] = db.get("command")
            out.append(info)
    return out


def pivot_info(pivot) -> dict:
    """Summarize an openpyxl TableDefinition into rows, columns, values, filters."""
    cache = pivot.cache
    fields = [f.name for f in cache.cacheFields]

    def names(idx_list):
        out = []
        for i in idx_list:
            try:
                out.append(fields[i])
            except (IndexError, TypeError):
                out.append(f"field_{i}")
        return out

    rows = names([f.x for f in pivot.rowFields])
    cols = names([f.x for f in pivot.colFields])
    filters = names([f.fld for f in pivot.pageFields])
    data = []
    for df in pivot.dataFields:
        try:
            fname = fields[df.fld]
        except (IndexError, TypeError):
            fname = f"field_{df.fld}"
        data.append({"field": fname, "agg": df.subtotal or "sum", "name": df.name})
    src = cache.cacheSource
    src_desc = None
    if src and src.worksheetSource is not None:
        ws_src = src.worksheetSource
        src_desc = {"sheet": ws_src.sheet, "ref": ws_src.ref, "table": ws_src.name}
    return {
        "name": pivot.name,
        "location": pivot.location.ref if pivot.location else None,
        "rows": rows,
        "cols": cols,
        "filters": filters,
        "values": data,
        "source": src_desc,
    }


def chart_info(chart) -> dict:
    title = None
    try:
        if chart.title and chart.title.tx and chart.title.tx.rich:
            title = "".join(
                r.t or "" for p in chart.title.tx.rich.p if p.r for r in p.r
            )
    except Exception:
        pass
    series = []
    for s in getattr(chart, "series", []) or []:
        try:
            ref = s.val.numRef.f if s.val and s.val.numRef else None
        except Exception:
            ref = None
        series.append(ref)
    return {"type": type(chart).__name__, "title": title, "series": series}


def dump(path: str, max_formulas: int) -> dict:
    wb = load_workbook(path, data_only=False)
    report: dict = {"file": path, "sheets": [], "defined_names": [], "m_queries": []}

    for dn_name, dn in wb.defined_names.items():
        report["defined_names"].append({"name": dn_name, "refers_to": dn.attr_text})

    for ws in wb.worksheets:
        info: dict = {
            "name": ws.title,
            "state": ws.sheet_state,
            "dims": ws.dimensions,
            "rows": ws.max_row,
            "cols": ws.max_column,
        }
        if ws.auto_filter and ws.auto_filter.ref:
            info["autofilter"] = ws.auto_filter.ref
        merged = [str(r) for r in ws.merged_cells.ranges]
        if merged:
            info["merged"] = merged[:20] + (
                [f"... +{len(merged) - 20} more"] if len(merged) > 20 else []
            )
        tables = []
        for t in getattr(ws, "tables", {}).values():
            tables.append(
                {
                    "name": t.name,
                    "ref": t.ref,
                    "columns": [c.name for c in t.tableColumns],
                }
            )
        if tables:
            info["tables"] = tables
        pivots = [pivot_info(p) for p in getattr(ws, "_pivots", [])]
        if pivots:
            info["pivots"] = pivots
        charts = [chart_info(c) for c in getattr(ws, "_charts", [])]
        if charts:
            info["charts"] = charts

        formula_counts: Counter = Counter()
        formula_example: dict = {}
        header_row = None
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    key = (col_of(cell.coordinate), normalize_formula(v))
                    formula_counts[key] += 1
                    formula_example.setdefault(key, (cell.coordinate, v))
            if header_row is None and any(c.value is not None for c in row):
                header_row = [
                    str(c.value) if c.value is not None else "" for c in row[:25]
                ]
        if header_row:
            info["first_nonempty_row"] = header_row
        if formula_counts:
            info["formulas"] = [
                {
                    "col": col,
                    "shape": shape,
                    "count": n,
                    "example_cell": formula_example[(col, shape)][0],
                    "example": formula_example[(col, shape)][1][:200],
                }
                for (col, shape), n in formula_counts.most_common(max_formulas)
            ]
        report["sheets"].append(info)

    report["m_queries"] = extract_m_queries(path)
    report["connections"] = extract_connections(path)
    return report


def render_text(r: dict) -> str:
    lines = [f"# {r['file']}"]
    for s in r["sheets"]:
        lines.append(
            f"\n## sheet '{s['name']}' ({s['state']}) dims={s['dims']} "
            f"rows={s['rows']} cols={s['cols']}"
        )
        if "first_nonempty_row" in s:
            lines.append(f"   header? {s['first_nonempty_row']}")
        if "autofilter" in s:
            lines.append(f"   autofilter: {s['autofilter']}")
        if "merged" in s:
            lines.append(f"   merged: {', '.join(s['merged'])}")
        for t in s.get("tables", []):
            lines.append(
                f"   table {t['name']} {t['ref']} cols={len(t['columns'])}: "
                + ", ".join(t["columns"][:12])
                + (" ..." if len(t["columns"]) > 12 else "")
            )
        for p in s.get("pivots", []):
            src = p["source"] or {}
            lines.append(
                f"   pivot {p['name']} @ {p['location']} "
                f"src={src.get('sheet')}!{src.get('ref') or src.get('table')}"
            )
            lines.append(
                f"      rows={p['rows']} cols={p['cols']} "
                f"filters={p['filters']} values={p['values']}"
            )
        for c in s.get("charts", []):
            lines.append(
                f"   chart {c['type']} title={c['title']!r} series={c['series']}"
            )
        for f in s.get("formulas", []):
            lines.append(
                f"   formula col {f['col']} x{f['count']} e.g. {f['example_cell']}: "
                f"{f['example']}"
            )
    if r.get("defined_names"):
        lines.append("\n## defined names")
        for dn in r["defined_names"]:
            lines.append(f"   {dn['name']} -> {dn['refers_to']}")
    if r.get("connections"):
        lines.append("\n## external connections")
        for c in r["connections"]:
            lines.append(
                f"   {c.get('name')}: {c.get('connection') or c.get('command') or c}"
            )
    if r.get("m_queries"):
        lines.append("\n## power query m")
        for q in r["m_queries"]:
            lines.append(f"--- {q['part']} :: {q['file']} ---")
            lines.append(q["code"].rstrip())
    return "\n".join(lines)


def col_idx(col_letters: str) -> int:
    n = 0
    for ch in col_letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def table_bounds(ws_struct, header_row: int | None) -> tuple[int, int, int] | None:
    for table in getattr(ws_struct, "tables", {}).values():
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", table.ref)
        if not m:
            continue
        start_col = col_idx(m.group(1))
        start_row = int(m.group(2))
        end_col = col_idx(m.group(3))
        end_row = int(m.group(4))
        if header_row is None or start_row == header_row:
            return start_row, start_col, end_col if end_col >= start_col else start_col
        if start_row <= header_row <= end_row:
            return header_row, start_col, end_col if end_col >= start_col else start_col
    return None


def detect_header_row(ws_vals) -> int:
    for row in range(1, ws_vals.max_row + 1):
        if any(
            ws_vals.cell(row, col).value not in (None, "")
            for col in range(1, ws_vals.max_column + 1)
        ):
            return row
    sys.exit("No non-empty rows found; cannot infer schema.")


def header_bounds(ws_vals, row_idx: int) -> tuple[int, int]:
    nonempty = [
        col
        for col in range(1, ws_vals.max_column + 1)
        if ws_vals.cell(row_idx, col).value not in (None, "")
    ]
    if not nonempty:
        sys.exit(f"Header row {row_idx} is empty; cannot infer schema.")
    return nonempty[0], nonempty[-1]


def infer_type(values: list[object]) -> str:
    if not values:
        return "string"
    if all(isinstance(v, bool) for v in values):
        return "boolean"
    if all(isinstance(v, (datetime, date)) for v in values):
        return "date"
    if all(isinstance(v, int) and not isinstance(v, bool) for v in values):
        return "bigint"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in values):
        return "double"
    return "string"


def infer_schema(path: str, sheet_name: str, header_row: int | None) -> dict:
    wb_vals = load_workbook(path, data_only=True, read_only=False)
    ws_vals = wb_vals[sheet_name]
    wb_struct = load_workbook(path, data_only=False, read_only=False)
    ws_struct = wb_struct[sheet_name]

    resolved_header_row = header_row or detect_header_row(ws_vals)
    bounds = table_bounds(ws_struct, resolved_header_row)
    if bounds is None:
        start_col, end_col = header_bounds(ws_vals, resolved_header_row)
    else:
        _, start_col, end_col = bounds

    columns = []
    for col in range(start_col, end_col + 1):
        header = ws_vals.cell(resolved_header_row, col).value
        name = str(header) if header not in (None, "") else get_column_letter(col)
        values = [
            ws_vals.cell(row, col).value
            for row in range(resolved_header_row + 1, ws_vals.max_row + 1)
            if ws_vals.cell(row, col).value is not None
        ]
        columns.append({"name": name, "type": infer_type(values)})
    return {"columns": columns, "userModified": True}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--max-formulas", type=int, default=20)
    ap.add_argument(
        "--schema",
        help="Emit Dataiku schema JSON from stored cells on a sheet.",
    )
    ap.add_argument("--header-row", type=int, help="1-based header row for --schema.")
    args = ap.parse_args()

    if args.schema:
        print(
            json.dumps(infer_schema(args.xlsx, args.schema, args.header_row), indent=2)
        )
        return

    report = dump(args.xlsx, args.max_formulas)
    if args.json:
        print(json.dumps(report, indent=2))
    else:
        print(render_text(report))


if __name__ == "__main__":
    main()

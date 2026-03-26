#!/usr/bin/env python3
"""
Parse the Wealth Manager FP&A Excel workbook into 6 clean CSVs for Dataiku upload.

Input:  main.xlsx (from ../test/brambles/)
Output: scripts/output/*.csv

Tabs parsed:
  1. Input Data       → input_data.csv          (117 records, unpivoted to long format)
  2. A - Asset Growth  → lookup_a_asset_growth.csv (17 rows × 4 years = 68 rows)
  3. B - Platform Wt   → lookup_b_platform_weight.csv (2 rows × 7 half-years = 14 rows)
  4. C - Share of Soln  → lookup_c_share_of_solution.csv (117 rows × 4 years = 468 rows)
  5. D - Seasonality    → lookup_d_seasonality.csv (8 rows × 12 months = 96 rows)
  6. E - Headcount      → lookup_e_headcount.csv  (1 segment × 48 months)
"""

import os
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

EXCEL_PATH = Path(__file__).resolve().parent.parent.parent / "test" / "brambles" / "main.xlsx"
OUTPUT_DIR = Path(__file__).resolve().parent / "output"


def parse_input_data(wb):
    """Parse Input Data tab: wide → long format.

    Wide format: 117 rows × 48 month columns (Jan 2025 – Dec 2028)
    Long format: record_no, dimensions..., month (YYYY-MM), value, act_or_plan
    """
    ws = wb["Input Data"]

    # Row 3 has Act/Plan labels (cols 16-63 = months, cols 64-69 = H1/H2/FY summaries)
    row3 = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]

    # Row 4 has column headers + month dates
    row4 = list(ws.iter_rows(min_row=4, max_row=4, values_only=True))[0]

    # Month columns: indices 16–63 (0-based), each has a datetime value
    month_cols = []
    for i in range(16, 64):
        dt = row4[i]
        if isinstance(dt, datetime):
            act_plan = row3[i] if row3[i] else "Plan"
            month_cols.append((i, dt.strftime("%Y-%m"), act_plan))

    # Dimension columns (fixed per record)
    dim_names = [
        "record_type",     # col 0 — always "New Business"
        "record_no",       # col 1
        "platform",        # col 2
        "channel",         # col 3
        "bp_channel",      # col 4
        "bp_segment",      # col 5
        "qi_switching",    # col 6
        "restricted",      # col 7
        "product",         # col 8
        "solution",        # col 9
        "admin_managed",   # col 10
        "core_noncore",    # col 11
        "fund_entity",     # col 12
        "growth_entity",   # col 13
        "concat_key_cf",   # col 14 — "Data with Cash Flow" concat key
        "concat_key",      # col 15 — "Data w/o Cash Flow" concat key (used for C join)
    ]

    rows = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] is None:
            continue
        dims = {dim_names[i]: row[i] for i in range(16)}
        for col_idx, month_str, act_plan in month_cols:
            val = row[col_idx]
            if val is None:
                val = 0
            # Derive join keys from month (YYYY-MM)
            year = int(month_str[:4])
            month_num = int(month_str[5:7])
            half_year = f"{year}_H1" if month_num <= 6 else f"{year}_H2"
            record = {
                **dims,
                "month": month_str,
                "year": year,
                "month_num": month_num,
                "half_year": half_year,
                "value": val,
                "act_or_plan": act_plan,
            }
            rows.append(record)

    return rows, ["record_type", "record_no", "platform", "channel", "bp_channel",
                   "bp_segment", "qi_switching", "restricted", "product", "solution",
                   "admin_managed", "core_noncore", "fund_entity", "growth_entity",
                   "concat_key_cf", "concat_key", "month", "year", "month_num",
                   "half_year", "value", "act_or_plan"]


def parse_tab_a(wb):
    """Parse Tab A - Asset Price Growth.

    4 year blocks side by side. Each block: key, company, solution, opening,
    annual_default, monthly_default, goal_seek_rate, calculated_rate, asset_price_growth.
    Output: growth_entity, solution, year, monthly_rate, goal_seek_rate, calculated_rate
    """
    ws = wb["A - Asset Price Growth"]

    # Year blocks start at cols: 0 (2025), 10 (2026), 20 (2027), 30 (2028)
    year_blocks = [
        (2025, 0),   # cols 0-8
        (2026, 10),  # cols 10-18
        (2027, 20),  # cols 20-28
        (2028, 30),  # cols 30-38
    ]

    rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        for year, offset in year_blocks:
            growth_entity = row[offset + 1]
            solution = row[offset + 2]
            monthly_default = row[offset + 5]
            goal_seek_rate = row[offset + 6]
            calculated_rate = row[offset + 7]
            if growth_entity and solution:
                rows.append({
                    "growth_entity": growth_entity,
                    "solution": solution,
                    "year": year,
                    "monthly_rate": monthly_default,
                    "goal_seek_rate": goal_seek_rate,
                    "calculated_rate": calculated_rate,
                })

    return rows, ["growth_entity", "solution", "year", "monthly_rate",
                   "goal_seek_rate", "calculated_rate"]


def parse_tab_b(wb):
    """Parse Tab B - Platform Weighting.

    Rows 8-9: Segment One × Platform 1/2, with weights per half-year.
    Output: bp_segment, platform, half_year, weight
    """
    ws = wb["B - Platform Weighting"]

    # Row 7 has headers: BP Segment, Platform, ..., 2025 H2, 2026 H1, 2026 H2, etc.
    half_years = [
        (4, "2025_H2"),
        (5, "2026_H1"),
        (6, "2026_H2"),
        (7, "2027_H1"),
        (8, "2027_H2"),
        (9, "2028_H1"),
        (10, "2028_H2"),
    ]

    rows = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if row[0] is None:
            continue
        bp_segment = row[0]
        platform = row[1]
        for col_idx, hy_label in half_years:
            weight = row[col_idx]
            if weight is not None:
                rows.append({
                    "bp_segment": bp_segment,
                    "platform": platform,
                    "half_year": hy_label,
                    "weight": weight,
                })

    return rows, ["bp_segment", "platform", "half_year", "weight"]


def parse_tab_c(wb):
    """Parse Tab C - Share of Solution.

    Already nearly clean. Cols: concat_key, platform, bp_segment, solution, product,
    entity, then 4 year columns (2025-2028).
    Output: concat_key, platform, bp_segment, solution, product, entity, year, share_rate
    """
    ws = wb["C- Share of Solution"]

    rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        concat_key = row[0]
        platform = row[1]
        bp_segment = row[2]
        solution = row[3]
        product = row[4]
        entity = row[5]
        for year_offset, year in enumerate([2025, 2026, 2027, 2028]):
            share_rate = row[6 + year_offset]
            if share_rate is None:
                share_rate = 0
            rows.append({
                "concat_key": concat_key,
                "platform": platform,
                "bp_segment": bp_segment,
                "solution": solution,
                "product": product,
                "entity": entity,
                "year": year,
                "share_rate": share_rate,
            })

    return rows, ["concat_key", "platform", "bp_segment", "solution", "product",
                   "entity", "year", "share_rate"]


def parse_tab_d(wb):
    """Parse Tab D - Seasonality.

    Channel 1 rows 9-12 (years 2025-2028), Channel 2 rows 16-19 (years 2025-2028).
    Each row has 12 monthly factors (cols 2-13).
    Output: channel, year, month_num, seasonality_factor
    """
    ws = wb["D - Seasonality"]

    month_names = ["January", "February", "March", "April", "May", "June",
                   "July", "August", "September", "October", "November", "December"]

    # Channel 1: rows 9-12, Channel 2: rows 16-19
    channel_rows = [
        ("Channel 1", 9, 12),
        ("Channel 2", 16, 19),
    ]

    rows = []
    for channel, start_row, end_row in channel_rows:
        for row in ws.iter_rows(min_row=start_row, max_row=end_row, values_only=True):
            if row[1] is None:
                continue
            year = int(row[1])
            for month_num in range(1, 13):
                factor = row[1 + month_num]  # cols 2-13
                if factor is None:
                    factor = 1.0
                rows.append({
                    "channel": channel,
                    "year": year,
                    "month_num": month_num,
                    "seasonality_factor": factor,
                })

    return rows, ["channel", "year", "month_num", "seasonality_factor"]


def parse_tab_e(wb):
    """Parse Tab E - Headcount & Productivity.

    MoM % Movement for Segment One is in row 56.
    Month labels in row 18 (cols 7-54 = Jan 2025 to Dec 2028).
    But MoM starts at col 8 (Feb 2025) since Jan has no prior month.

    We also need Jan 2025 = 1.0 as the baseline (no change from prior month total).
    For the forecast, we use Jul 2025 onward. But we include all months for completeness.

    Output: bp_segment, month (YYYY-MM), headcount_mom
    """
    ws = wb["E - Headcount & productivity"]

    # Row 18: month labels (col 7 = Jan 2025, col 8 = Feb 2025, ..., col 54 = Dec 2028)
    row18 = list(ws.iter_rows(min_row=18, max_row=18, values_only=True))[0]

    # Row 56: MoM values (col 7 = None/Jan, col 8 = Feb 2025 MoM, ...)
    row56 = list(ws.iter_rows(min_row=56, max_row=56, values_only=True))[0]

    rows = []
    for col in range(7, 55):
        month_label = row18[col]
        if month_label is None:
            continue
        # Parse "Jan 2025" → "2025-01"
        try:
            dt = datetime.strptime(str(month_label), "%b %Y")
            month_str = dt.strftime("%Y-%m")
        except ValueError:
            continue

        mom = row56[col] if col < len(row56) and row56[col] is not None else 1.0
        rows.append({
            "bp_segment": "Segment One",
            "month": month_str,
            "headcount_mom": mom,
        })

    return rows, ["bp_segment", "month", "headcount_mom"]


def write_csv(rows, columns, filename):
    """Write rows to CSV."""
    filepath = OUTPUT_DIR / filename
    with open(filepath, "w") as f:
        f.write(",".join(columns) + "\n")
        for row in rows:
            vals = []
            for col in columns:
                v = row.get(col, "")
                if isinstance(v, str) and ("," in v or '"' in v or "\n" in v):
                    v = '"' + v.replace('"', '""') + '"'
                else:
                    v = str(v) if v is not None else ""
                vals.append(v)
            f.write(",".join(vals) + "\n")
    print(f"  {filename}: {len(rows)} rows")


def main():
    if not EXCEL_PATH.exists():
        print(f"ERROR: Excel file not found at {EXCEL_PATH}")
        sys.exit(1)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print(f"Reading {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)

    print("Parsing sheets...")

    # 1. Input Data (wide → long)
    rows, cols = parse_input_data(wb)
    write_csv(rows, cols, "input_data.csv")

    # 2. Tab A - Asset Price Growth
    rows, cols = parse_tab_a(wb)
    write_csv(rows, cols, "lookup_a_asset_growth.csv")

    # 3. Tab B - Platform Weighting
    rows, cols = parse_tab_b(wb)
    write_csv(rows, cols, "lookup_b_platform_weight.csv")

    # 4. Tab C - Share of Solution
    rows, cols = parse_tab_c(wb)
    write_csv(rows, cols, "lookup_c_share_of_solution.csv")

    # 5. Tab D - Seasonality
    rows, cols = parse_tab_d(wb)
    write_csv(rows, cols, "lookup_d_seasonality.csv")

    # 6. Tab E - Headcount & Productivity
    rows, cols = parse_tab_e(wb)
    write_csv(rows, cols, "lookup_e_headcount.csv")

    print("\nDone! CSVs written to", OUTPUT_DIR)


if __name__ == "__main__":
    main()

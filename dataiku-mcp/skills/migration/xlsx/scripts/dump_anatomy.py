#!/usr/bin/env python3
"""Dump the migration-relevant anatomy of an .xlsx workbook.

Phase-1 inventory helper for Excel -> Dataiku migrations. Reports, per sheet:
dimensions, tables (ListObjects), pivot tables (rows/cols/values/filters and
their source range), charts, distinct formulas (deduped by R1C1 shape with
counts), merged cells, autofilters -- plus workbook-level named ranges,
defined names, external connections, and embedded Power Query (M) source
when present.

Usage:
    uv run --with openpyxl python dump_anatomy.py BOOK.xlsx [--json] [--max-formulas N]
    uv run --with openpyxl python dump_anatomy.py BOOK.xlsx --schema "Sheet Name"

The M extraction reads customXml/item*.xml DataMashup blobs (base64 -> inner
ZIP -> Formulas/Section1.m). No external deps beyond openpyxl.
"""

from __future__ import annotations

import argparse
import base64
import io
import json
import re
import sys
import zipfile
from collections import Counter

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def col_of(cell_ref: str) -> str:
    m = re.match(r"([A-Z]+)", cell_ref)
    return m.group(1) if m else ""


def normalize_formula(f: str) -> str:
    """Collapse row numbers so per-row copies of one formula dedupe together."""
    return re.sub(r"(?<![A-Za-z0-9_])(\$?[A-Z]{1,3})\$?\d+", r"\1#", f)


def extract_m_queries(path: str) -> list[dict]:
    """Pull Power Query M section files out of the DataMashup customXml part.

    customXml/item*.xml is typically UTF-16 (BOM-prefixed) -- decode before
    regexing, a bytes-level search silently misses the tag.
    """
    out = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not re.match(r"customXml/item\d+\.xml$", name):
                continue
            raw = z.read(name)
            enc = "utf-16" if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8"
            data = raw.decode(enc, "replace")
            m = re.search(r"<DataMashup[^>]*>(.*?)</DataMashup>", data, re.S)
            if not m:
                continue
            try:
                blob = base64.b64decode(m.group(1))
            except Exception:
                continue
            # MS-QDEFF: [version:4][package_len:4][package ZIP][permissions]
            # [metadata]... Slice the package ZIP by declared length -- ZipFile
            # on the whole blob finds the EOCD of trailing sections instead.
            if len(blob) < 8:
                continue
            pkg_len = int.from_bytes(blob[4:8], "little")
            pkg = blob[8 : 8 + pkg_len]
            if not pkg.startswith(b"PK\x03\x04"):
                pk = blob.find(b"PK\x03\x04")
                if pk < 0:
                    continue
                pkg = blob[pk:]
            try:
                inner = zipfile.ZipFile(io.BytesIO(pkg))
            except Exception:
                continue
            for iname in inner.namelist():
                if iname.endswith(".m"):
                    out.append(
                        {
                            "part": name,
                            "file": iname,
                            "code": inner.read(iname).decode("utf-8", "replace"),
                        }
                    )
    return out


def extract_connections(path: str) -> list[dict]:
    """Summarize xl/connections.xml -- external data sources behind queryTables."""
    import xml.etree.ElementTree as ET

    out = []
    with zipfile.ZipFile(path) as z:
        if "xl/connections.xml" not in z.namelist():
            return out
        root = ET.fromstring(z.read("xl/connections.xml"))
        ns = {"m": root.tag.split("}")[0].strip("{")} if "}" in root.tag else {}
        conns = root.findall("m:connection", ns) if ns else root.findall("connection")
        for c in conns:
            info = {
                "name": c.get("name"),
                "description": c.get("description"),
                "type": c.get("type"),
            }
            db = c.find("m:dbPr", ns) if ns else c.find("dbPr")
            if db is not None:
                info["connection"] = db.get("connection")
                info["command"] = db.get("command")
            out.append(info)
    return out


def pivot_info(pivot) -> dict:
    """Summarize an openpyxl TableDefinition into rows/cols/values/filters."""
    cache = pivot.cache
    fields = [f.name for f in cache.cacheFields]

    def names(idx_list):
        out = []
        for i in idx_list:
            try:
                out.append(fields[i])
            except (IndexError, TypeError):
                out.append(f"field_{i}")
        return out

    rows = names([f.x for f in pivot.rowFields])
    cols = names([f.x for f in pivot.colFields])
    filters = names([f.fld for f in pivot.pageFields])
    data = []
    for df in pivot.dataFields:
        try:
            fname = fields[df.fld]
        except (IndexError, TypeError):
            fname = f"field_{df.fld}"
        data.append({"field": fname, "agg": df.subtotal or "sum", "name": df.name})
    src = cache.cacheSource
    src_desc = None
    if src and src.worksheetSource is not None:
        ws_src = src.worksheetSource
        src_desc = {"sheet": ws_src.sheet, "ref": ws_src.ref, "table": ws_src.name}
    return {
        "name": pivot.name,
        "location": pivot.location.ref if pivot.location else None,
        "rows": rows,
        "cols": cols,
        "filters": filters,
        "values": data,
        "source": src_desc,
    }


def chart_info(chart) -> dict:
    title = None
    try:
        if chart.title and chart.title.tx and chart.title.tx.rich:
            title = "".join(
                r.t or "" for p in chart.title.tx.rich.p if p.r for r in p.r
            )
    except Exception:
        pass
    series = []
    for s in getattr(chart, "series", []) or []:
        try:
            ref = s.val.numRef.f if s.val and s.val.numRef else None
        except Exception:
            ref = None
        series.append(ref)
    return {"type": type(chart).__name__, "title": title, "series": series}


def dump(path: str, max_formulas: int) -> dict:
    wb = load_workbook(path, data_only=False)
    report: dict = {"file": path, "sheets": [], "defined_names": [], "m_queries": []}

    for dn_name, dn in wb.defined_names.items():
        report["defined_names"].append({"name": dn_name, "refers_to": dn.attr_text})

    for ws in wb.worksheets:
        info: dict = {
            "name": ws.title,
            "state": ws.sheet_state,
            "dims": ws.dimensions,
            "rows": ws.max_row,
            "cols": ws.max_column,
        }
        if ws.auto_filter and ws.auto_filter.ref:
            info["autofilter"] = ws.auto_filter.ref
        merged = [str(r) for r in ws.merged_cells.ranges]
        if merged:
            info["merged"] = merged[:20] + (
                [f"... +{len(merged) - 20} more"] if len(merged) > 20 else []
            )
        tables = []
        for t in getattr(ws, "tables", {}).values():
            tables.append(
                {
                    "name": t.name,
                    "ref": t.ref,
                    "columns": [c.name for c in t.tableColumns],
                }
            )
        if tables:
            info["tables"] = tables
        pivots = [pivot_info(p) for p in getattr(ws, "_pivots", [])]
        if pivots:
            info["pivots"] = pivots
        charts = [chart_info(c) for c in getattr(ws, "_charts", [])]
        if charts:
            info["charts"] = charts

        formula_counts: Counter = Counter()
        formula_example: dict = {}
        header_row = None
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    key = (col_of(cell.coordinate), normalize_formula(v))
                    formula_counts[key] += 1
                    formula_example.setdefault(key, (cell.coordinate, v))
            if header_row is None and any(c.value is not None for c in row):
                header_row = [
                    str(c.value) if c.value is not None else "" for c in row[:25]
                ]
        if header_row:
            info["first_nonempty_row"] = header_row
        if formula_counts:
            info["formulas"] = [
                {
                    "col": col,
                    "shape": shape,
                    "count": n,
                    "example_cell": formula_example[(col, shape)][0],
                    "example": formula_example[(col, shape)][1][:200],
                }
                for (col, shape), n in formula_counts.most_common(max_formulas)
            ]
        report["sheets"].append(info)

    report["m_queries"] = extract_m_queries(path)
    report["connections"] = extract_connections(path)
    return report


def render_text(r: dict) -> str:
    lines = [f"# {r['file']}"]
    for s in r["sheets"]:
        lines.append(
            f"\n## sheet '{s['name']}' ({s['state']}) dims={s['dims']} "
            f"rows={s['rows']} cols={s['cols']}"
        )
        if "first_nonempty_row" in s:
            lines.append(f"   header? {s['first_nonempty_row']}")
        if "autofilter" in s:
            lines.append(f"   autofilter: {s['autofilter']}")
        if "merged" in s:
            lines.append(f"   merged: {s['merged']}")
        for t in s.get("tables", []):
            lines.append(f"   table {t['name']} {t['ref']} cols={t['columns']}")
        for p in s.get("pivots", []):
            lines.append(
                f"   pivot {p['name']} @ {p['location']} src={p['source']}\n"
                f"      rows={p['rows']} cols={p['cols']} filters={p['filters']}\n"
                f"      values={p['values']}"
            )
        for c in s.get("charts", []):
            lines.append(
                f"   chart {c['type']} title={c['title']!r} series={c['series']}"
            )
        for f in s.get("formulas", []):
            lines.append(
                f"   formula col {f['col']} x{f['count']} "
                f"@{f['example_cell']}: {f['example']}"
            )
    if r["defined_names"]:
        lines.append("\n## defined names")
        for d in r["defined_names"]:
            lines.append(f"   {d['name']} = {d['refers_to']}")
    if r.get("connections"):
        lines.append("\n## connections (xl/connections.xml)")
        for c in r["connections"]:
            lines.append(f"   {json.dumps(c, default=str)}")
    for q in r["m_queries"]:
        lines.append(f"\n## M query {q['part']}::{q['file']}\n{q['code']}")
    return "\n".join(lines)


def emit_dss_schema(path: str, sheet: str, header_row: int = 1) -> list[dict]:
    """Infer a DSS schema (name/type) for one sheet from stored cell types.

    Scans data rows below header_row: all-int -> bigint, any float ->
    double, datetime -> date, else string. Columns are clamped to the
    header extent (used-range junk to the right is dropped).
    """
    import datetime as dt

    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    headers: list[tuple[int, str]] = []
    for idx, cell in enumerate(
        next(ws.iter_rows(min_row=header_row, max_row=header_row)), 1
    ):
        if cell.value is not None and str(cell.value).strip():
            headers.append((idx, str(cell.value).strip()))
    cols = []
    for idx, name in headers:
        letter = get_column_letter(idx)
        seen_float = seen_int = seen_date = seen_str = False
        n = 0
        for (v,) in ws.iter_rows(
            min_row=header_row + 1, min_col=idx, max_col=idx, values_only=True
        ):
            if v is None:
                continue
            n += 1
            if isinstance(v, bool) or isinstance(v, str):
                seen_str = True
            elif isinstance(v, dt.datetime) or isinstance(v, dt.date):
                seen_date = True
            elif isinstance(v, float):
                if v.is_integer():
                    seen_int = True
                else:
                    seen_float = True
            elif isinstance(v, int):
                seen_int = True
            else:
                seen_str = True
            if n >= 5000:
                break
        if seen_str or n == 0:
            typ = "string"
        elif seen_date and not (seen_float or seen_int):
            typ = "date"
        elif seen_float:
            typ = "double"
        elif seen_int:
            typ = "bigint"
        else:
            typ = "string"
        cols.append({"name": name, "type": typ, "_excel_col": letter})
    return cols


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbook")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--max-formulas", type=int, default=15)
    ap.add_argument(
        "--schema",
        metavar="SHEET",
        help="Emit a DSS schema JSON (columns name/type) for one sheet and exit",
    )
    ap.add_argument(
        "--header-row", type=int, default=1, help="1-based header row for --schema"
    )
    args = ap.parse_args()
    if args.schema:
        cols = emit_dss_schema(args.workbook, args.schema, args.header_row)
        for c in cols:
            c.pop("_excel_col", None)
        json.dump({"columns": cols, "userModified": True}, sys.stdout, indent=1)
        return
    r = dump(args.workbook, args.max_formulas)
    if args.json:
        json.dump(r, sys.stdout, indent=2, default=str)
    else:
        print(render_text(r))


if __name__ == "__main__":
    main()

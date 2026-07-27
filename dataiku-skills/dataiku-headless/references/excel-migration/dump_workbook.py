#!/usr/bin/env python3
"""Dump workbook cells with formulas, array extents, and cached values · in: <workbook.xlsx> <outdir> → out: one TSV per sheet + stdout digest · deps: openpyxl.

Loads the workbook twice (data_only=False and True) and writes one TSV per
sheet:

    COORD <tab> F|A|V <tab> formula-or-value [<tab> => <tab> cached-value]

F = formula, A = array formula, V = literal value. Also prints a per-sheet
summary and lists defined names that still resolve.

Usage:
    python dump_workbook.py BOOK.xlsx OUTDIR [SHEET ...]
"""

import sys
import warnings
from pathlib import Path

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

warnings.filterwarnings("ignore")


def main() -> None:
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    book, outdir = sys.argv[1], Path(sys.argv[2])
    only = set(sys.argv[3:])
    outdir.mkdir(parents=True, exist_ok=True)

    wbf = openpyxl.load_workbook(book, data_only=False)
    wbv = openpyxl.load_workbook(book, data_only=True)

    broken = sum(
        1
        for dn in wbf.defined_names.values()
        if dn.attr_text and "#REF!" in dn.attr_text
    )
    live = [
        (n, dn.attr_text)
        for n, dn in wbf.defined_names.items()
        if dn.attr_text and "#REF!" not in dn.attr_text
    ]
    print(f"defined names: {len(live)} live, {broken} broken (#REF!, ignore)")
    for n, t in live[:30]:
        print(f"  {n} -> {t[:100]}")

    for ws in wbf.worksheets:
        if only and ws.title not in only:
            continue
        wsv = wbv[ws.title]
        lines, n_f, n_a = [], 0, 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                cached = wsv.cell(row=c.row, column=c.column).value
                v = c.value
                if isinstance(v, ArrayFormula):
                    n_a += 1
                    lines.append(
                        f"{c.coordinate}\tA\t[{v.ref}] "
                        f"{str(v.text).replace(chr(10), ' ')}\t=>\t{cached!r}"
                    )
                elif isinstance(v, str) and v.startswith("="):
                    n_f += 1
                    lines.append(
                        f"{c.coordinate}\tF\t{v.replace(chr(10), ' ')}\t=>\t{cached!r}"
                    )
                else:
                    lines.append(f"{c.coordinate}\tV\t{v!r}")
        safe = ws.title.replace(" ", "_").replace("/", "_").replace(">", "")
        path = outdir / f"{safe}.tsv"
        path.write_text("\n".join(lines))
        role = (
            "engine" if n_f > len(lines) * 0.4 else ("input" if n_f == 0 else "mixed")
        )
        print(
            f"{ws.title}: {len(lines)} cells ({n_f} formulas, {n_a} array, "
            f"{ws.max_row}x{ws.max_column}, state={ws.sheet_state}) "
            f"-> {path}  [{role}?]"
        )


if __name__ == "__main__":
    main()
